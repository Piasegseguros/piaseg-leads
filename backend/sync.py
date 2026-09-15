import io
import json
import os
from datetime import datetime, timezone

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

from database import SessionLocal
from models import Lead

SHEET_FILE_NAME = "CRM SEGURO JÁ.xlsx"
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

TEST_EMAILS = {"test@meta.com"}


def _drive_service():
    raw = os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"]
    info = json.loads(raw)
    creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
    return build("drive", "v3", credentials=creds)


def _download_workbook() -> openpyxl.Workbook:
    service = _drive_service()
    results = service.files().list(
        q=f"name = '{SHEET_FILE_NAME}'",
        fields="files(id, name)",
        pageSize=1,
    ).execute()
    files = results.get("files", [])
    if not files:
        raise RuntimeError(f"Arquivo '{SHEET_FILE_NAME}' não encontrado ou não compartilhado com a conta de serviço")

    file_id = files[0]["id"]
    request = service.files().get_media(fileId=file_id)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buffer.seek(0)
    return openpyxl.load_workbook(buffer, read_only=True, data_only=True)


def _parse_rows(wb: openpyxl.Workbook):
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {name: i for i, name in enumerate(header)}

    for row in rows:
        if row is None or row[idx["id"]] is None:
            continue
        email = (row[idx.get("email")] or "").strip() if idx.get("email") is not None else ""
        full_name = (row[idx["full_name"]] or "").strip()
        if email in TEST_EMAILS or full_name.startswith("<"):
            continue

        phone = row[idx["phone_number"]] if idx.get("phone_number") is not None else None
        if phone:
            phone = phone.replace("p:", "")

        created_time = row[idx["created_time"]]
        if isinstance(created_time, str):
            created_time = datetime.fromisoformat(created_time)

        yield {
            "id": row[idx["id"]],
            "created_time": created_time,
            "full_name": full_name,
            "email": email or None,
            "phone_number": phone,
            "campaign_name": row[idx.get("campaign_name")] if idx.get("campaign_name") is not None else None,
            "ad_name": row[idx.get("ad_name")] if idx.get("ad_name") is not None else None,
        }


def sync_leads() -> int:
    wb = _download_workbook()
    new_count = 0
    db = SessionLocal()
    try:
        existing_ids = {lid for (lid,) in db.query(Lead.id).all()}
        now = datetime.now(timezone.utc)
        for data in _parse_rows(wb):
            if data["id"] in existing_ids:
                continue
            lead = Lead(synced_at=now, **data)
            db.add(lead)
            new_count += 1
        db.commit()
    finally:
        db.close()
    return new_count
